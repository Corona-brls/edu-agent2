"""Parse knowledge documents into common Document/Page objects."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_SUFFIXES = {".md", ".txt", ".pdf", ".docx"}
DEFAULT_MAX_FILE_BYTES = 25 * 1024 * 1024


class DocumentParseError(RuntimeError):
    """Raised when a document cannot be parsed into text."""


@dataclass(frozen=True)
class Page:
    text: str
    page: int | None
    section: str
    source_type: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class Document:
    source_path: Path
    filename: str
    file_type: str
    title: str
    text: str
    pages: list[Page] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def source_type(self) -> str:
        return self.file_type


def validate_upload_filename(filename: str | None) -> str:
    raw_name = (filename or "").strip()
    safe_name = Path(raw_name).name
    if not raw_name or safe_name != raw_name or safe_name in {".", ".."} or ".." in Path(raw_name).parts:
        raise DocumentParseError("文件名不合法，不能包含路径或上级目录。")
    return safe_name


def parse_document(path: str | Path, max_file_bytes: int = DEFAULT_MAX_FILE_BYTES) -> Document:
    source_path = Path(path)
    if max_file_bytes <= 0:
        raise ValueError("max_file_bytes 必须大于 0")
    if not source_path.exists():
        raise DocumentParseError(f"文档不存在：{source_path}")
    if not source_path.is_file():
        raise DocumentParseError(f"不是有效文件：{source_path}")
    if source_path.stat().st_size > max_file_bytes:
        raise DocumentParseError(f"文档超过大小限制：{source_path.name}，最大 {max_file_bytes} 字节")

    suffix = source_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise DocumentParseError(f"不支持的文档格式：{suffix or '无扩展名'}。支持格式：{supported}")

    file_type = suffix.lstrip(".")
    if suffix in {".md", ".txt"}:
        pages = _parse_text_file(source_path, file_type)
        title = _title_from_pages(pages, source_path.stem)
        metadata: dict[str, Any] = {"parser": "text", "source_type": file_type}
    elif suffix == ".pdf":
        pages, metadata = _parse_pdf(source_path)
        title = _title_from_pages(pages, source_path.stem)
    elif suffix == ".docx":
        pages, metadata = _parse_docx(source_path)
        title = metadata.get("title") or _title_from_pages(pages, source_path.stem)
    else:
        raise DocumentParseError(f"未实现的文档格式：{suffix}")

    pages = [page for page in pages if page.text.strip()]
    text = _normalize_text("\n\n".join(page.text for page in pages))
    if not text:
        raise DocumentParseError(f"未能从文档中提取到有效文本：{source_path.name}")

    metadata.update(
        {
            "filename": source_path.name,
            "file_type": file_type,
            "source_type": file_type,
            "page_count": len(pages),
            "byte_size": source_path.stat().st_size,
        }
    )
    return Document(
        source_path=source_path,
        filename=source_path.name,
        file_type=file_type,
        title=title,
        text=text,
        pages=pages,
        metadata=metadata,
    )


def chunk_document_text(text: str, max_chars: int = 800, overlap: int = 80) -> list[str]:
    if max_chars <= 0:
        raise ValueError("max_chars 必须大于 0")
    if overlap < 0 or overlap >= max_chars:
        raise ValueError("overlap 必须大于等于 0 且小于 max_chars")

    text = _normalize_text(text)
    if not text:
        return []

    blocks = [block.strip() for block in text.split("\n") if block.strip()]
    chunks: list[str] = []
    current = ""

    for block in blocks:
        candidate = f"{current}\n{block}".strip()
        if len(candidate) <= max_chars:
            current = candidate
            continue
        if current:
            chunks.append(current)
        current = block

        while len(current) > max_chars:
            chunks.append(current[:max_chars].strip())
            current = current[max(0, max_chars - overlap) :].strip()

    if current:
        chunks.append(current)

    return chunks


def _parse_text_file(path: Path, file_type: str) -> list[Page]:
    try:
        raw_text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise DocumentParseError(f"文本文件必须使用 UTF-8 编码：{path.name}") from exc
    except OSError as exc:
        raise DocumentParseError(f"读取文本文件失败：{path.name}") from exc

    section = _first_section(raw_text, path.stem)
    return [
        Page(
            text=_normalize_text(raw_text),
            page=1,
            section=section,
            source_type=file_type,
            metadata={"filename": path.name, "line_start": 1},
        )
    ]


def _parse_pdf(path: Path) -> tuple[list[Page], dict[str, Any]]:
    try:
        reader = PdfReader(str(path))
        pages: list[Page] = []
        for page_index, page in enumerate(reader.pages, start=1):
            text = _normalize_text(page.extract_text() or "")
            if text:
                pages.append(
                    Page(
                        text=text,
                        page=page_index,
                        section=f"page {page_index}",
                        source_type="pdf",
                        metadata={"filename": path.name},
                    )
                )
    except Exception as exc:  # pypdf raises multiple exception types for malformed PDFs.
        raise DocumentParseError(f"解析 PDF 失败：{path.name}") from exc

    return pages, {"parser": "pypdf", "pages": len(reader.pages), "source_type": "pdf"}


def _parse_docx(path: Path) -> tuple[list[Page], dict[str, Any]]:
    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentParseError(f"解析 DOCX 失败：{path.name}") from exc

    pages: list[Page] = []
    section = (doc.core_properties.title or "").strip() or path.stem
    current_parts: list[str] = []

    def flush_section() -> None:
        nonlocal current_parts
        text = _normalize_text("\n".join(current_parts))
        if text:
            pages.append(
                Page(
                    text=text,
                    page=None,
                    section=section,
                    source_type="docx",
                    metadata={"filename": path.name},
                )
            )
        current_parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            flush_section()
            section = text
            current_parts.append(text)
        else:
            current_parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                current_parts.append(" | ".join(cells))
    flush_section()

    metadata = {
        "parser": "python-docx",
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "title": (doc.core_properties.title or "").strip(),
        "source_type": "docx",
    }
    return pages, metadata


def _title_from_pages(pages: list[Page], fallback: str) -> str:
    for page in pages:
        title = _first_section(page.text, fallback)
        if title:
            return title
    return fallback


def _first_section(text: str, fallback: str) -> str:
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("#"):
            return line.lstrip("#").strip() or fallback
        return line[:80]
    return fallback


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line)
