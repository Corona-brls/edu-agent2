from pathlib import Path

import pytest
from docx import Document as DocxDocument

from backend.rag.document_parser import (
    DocumentParseError,
    chunk_document_text,
    parse_document,
    validate_upload_filename,
)


def test_parse_markdown(tmp_path):
    path = tmp_path / "sample.md"
    path.write_text("# Markdown 标题\n\n这是一段 Markdown 内容。", encoding="utf-8")

    document = parse_document(path)

    assert document.file_type == "md"
    assert document.title == "Markdown 标题"
    assert "Markdown 内容" in document.text
    assert document.metadata["source_type"] == "md"
    assert document.pages[0].page == 1
    assert document.pages[0].section == "Markdown 标题"
    assert document.pages[0].source_type == "md"


def test_parse_txt(tmp_path):
    path = tmp_path / "sample.txt"
    path.write_text("TXT 标题\n这是一段 TXT 内容。", encoding="utf-8")

    document = parse_document(path)

    assert document.file_type == "txt"
    assert document.title == "TXT 标题"
    assert "TXT 内容" in document.text
    assert document.metadata["source_type"] == "txt"
    assert document.pages[0].page == 1
    assert document.pages[0].section == "TXT 标题"
    assert document.pages[0].source_type == "txt"


def test_parse_pdf(tmp_path):
    path = tmp_path / "sample.pdf"
    _write_minimal_text_pdf(path, "PDF sample text")

    document = parse_document(path)

    assert document.file_type == "pdf"
    assert document.metadata["pages"] == 1
    assert "PDF sample text" in document.text
    assert document.metadata["source_type"] == "pdf"
    assert document.pages[0].page == 1
    assert document.pages[0].section == "page 1"
    assert document.pages[0].source_type == "pdf"


def test_parse_docx(tmp_path):
    path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.core_properties.title = "DOCX 标题"
    doc.add_heading("DOCX 标题", level=1)
    doc.add_paragraph("这是一段 DOCX 内容。")
    doc.save(path)

    document = parse_document(path)

    assert document.file_type == "docx"
    assert document.title == "DOCX 标题"
    assert "DOCX 内容" in document.text
    assert document.metadata["source_type"] == "docx"
    assert document.pages[0].page is None
    assert document.pages[0].section == "DOCX 标题"
    assert document.pages[0].source_type == "docx"


def test_parse_unsupported_file_has_clear_error(tmp_path):
    path = tmp_path / "sample.xlsx"
    path.write_text("not supported", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="不支持的文档格式"):
        parse_document(path)


def test_parse_missing_file_has_clear_error(tmp_path):
    with pytest.raises(DocumentParseError, match="文档不存在"):
        parse_document(tmp_path / "missing.pdf")


def test_parse_bad_utf8_has_clear_error(tmp_path):
    path = tmp_path / "bad.txt"
    path.write_bytes(b"\xff\xfe\x00")

    with pytest.raises(DocumentParseError, match="UTF-8"):
        parse_document(path)


def test_parse_empty_file_has_clear_error(tmp_path):
    path = tmp_path / "empty.md"
    path.write_text("", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="有效文本"):
        parse_document(path)


def test_parse_corrupt_pdf_has_clear_error(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"not a real pdf")

    with pytest.raises(DocumentParseError, match="解析 PDF 失败"):
        parse_document(path)


def test_parse_corrupt_docx_has_clear_error(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a real docx")

    with pytest.raises(DocumentParseError, match="解析 DOCX 失败"):
        parse_document(path)


def test_parse_rejects_oversized_file(tmp_path):
    path = tmp_path / "large.txt"
    path.write_text("too large", encoding="utf-8")

    with pytest.raises(DocumentParseError, match="大小限制"):
        parse_document(path, max_file_bytes=3)


def test_validate_upload_filename_rejects_path_traversal():
    assert validate_upload_filename("safe.md") == "safe.md"
    with pytest.raises(DocumentParseError, match="文件名不合法"):
        validate_upload_filename("../evil.md")
    with pytest.raises(DocumentParseError, match="文件名不合法"):
        validate_upload_filename("folder/evil.md")
    with pytest.raises(DocumentParseError, match="文件名不合法"):
        validate_upload_filename(r"folder\evil.md")


def test_chunk_document_text_rejects_invalid_overlap():
    with pytest.raises(ValueError, match="overlap"):
        chunk_document_text("long text", max_chars=10, overlap=10)


def test_chunk_document_text_rejects_invalid_max_chars():
    with pytest.raises(ValueError, match="max_chars"):
        chunk_document_text("long text", max_chars=0)

def _write_minimal_text_pdf(path: Path, text: str) -> None:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{index} 0 obj\n".encode("ascii"))
        content.extend(obj)
        content.extend(b"\nendobj\n")

    xref_offset = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    content.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(bytes(content))

